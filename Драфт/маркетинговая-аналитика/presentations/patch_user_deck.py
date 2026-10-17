# -*- coding: utf-8 -*-
"""Патч пользовательской версии деки «Лекция 3-4» (не регенерация!).

1) 13 вставленных слайдов из старой деки приведены к стилю X5
   (четыре схемы моделей склеены в один слайд 2×2);
2) добавлены «под капотом» ML-атрибуции: шаг 1 (матрица признаков) и шаг 2
   (обучение) после пайплайна vc.ru; картинка слайда «Вклад касания = разница
   прогнозов» заменена на шаг 3 — разбор юзера с реальными числами practice_4.3;
3) страницы перенумерованы.

Запуск: python3 presentations/patch_user_deck.py <путь к pptx>
"""
from __future__ import annotations
import copy
import hashlib
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from make_decks_x5 import clone_slide, fill_bullets, set_text, title_font_size

DRAFT = HERE.parent
HOOD = DRAFT / "course" / "modules" / "M4" / "images"
ASSETS = Path("/tmp/user_assets")
SLIDE_W = 20.0


# ---------------------------------------------------------------- утилиты
def boxes_by_id(slide):
    return {sh.shape_id: sh for sh in slide.shapes if sh.has_text_frame}


def slide_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def find_slide(prs, title_prefix):
    for s in prs.slides:
        if slide_title(s).startswith(title_prefix):
            return s
    raise KeyError(title_prefix)


def make_pic_slide(prs, tpl, title, img_path, caption):
    s = clone_slide(prs, tpl)
    b = boxes_by_id(s)
    set_text(b[5], title, size=title_font_size(title, 54))
    set_text(b[6], caption, size=20)
    for i in (7, 8, 9, 10, 11):
        set_text(b[i], "")
    w_px, h_px = Image.open(img_path).size
    max_h, max_w = 7.15, 18.4
    scale = min(max_h / (h_px / 150.0), max_w / (w_px / 150.0))
    pic_w, pic_h = w_px / 150.0 * scale, h_px / 150.0 * scale
    s.shapes.add_picture(str(img_path), Inches((SLIDE_W - pic_w) / 2),
                         Inches(2.75), width=Inches(pic_w), height=Inches(pic_h))
    return s


def make_bul_slide(prs, tpl, title, bullets, left_head):
    s = clone_slide(prs, tpl)
    # прямое заполнение по id носителя (в источнике боксы могут быть пустыми,
    # поэтому не пользуемся fill_bullets с фильтром по непустым)
    b = {sh.shape_id: sh for sh in s.shapes if sh.has_text_frame}
    set_text(b[5], title, size=title_font_size(title, 54))
    set_text(b[6], "", size=22)
    set_text(b[7], left_head, size=28)
    set_text(b[8], "", size=20)
    set_text(b[9], "", size=28)
    set_text(b[10], "", size=20)
    box = b[11]
    box.top, box.height = Inches(4.6), Inches(5.6)
    items = [x if len(x) <= 56 else x[:53] + "…" for x in bullets][:7]
    set_text(box, items, size=22)
    return s


def renumber(prs):
    for n, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().isdigit():
                set_text(sh, str(n))
                break


# ---------------------------------------------------------------- подготовка картинок
def prep_assets():
    # s25_5 — 5.9 МБ, пережмём
    im = Image.open(ASSETS / "s25_5.png").convert("RGB")
    if im.width > 1600:
        im = im.resize((1600, int(im.height * 1600 / im.width)), Image.LANCZOS)
    im.save(ASSETS / "s25_small.png")
    # композит 2×2 из четырёх схем моделей
    cells = [("s35_4.png", "first-click: всё первому"),
             ("s36_5.png", "last-click: всё последнему"),
             ("s37_6.png", "linear: поровну"),
             ("s38_7.png", "затухание к финалу (time-decay)")]
    thumbs = []
    for f, label in cells:
        im = Image.open(ASSETS / f).convert("RGB")
        im.thumbnail((1050, 900), Image.LANCZOS)
        thumbs.append((im, label))
    cw = max(im.width for im, _ in thumbs) + 24
    ch = max(im.height for im, _ in thumbs) + 46
    sheet = Image.new("RGB", (cw * 2 + 16, ch * 2 + 16), "white")
    d = ImageDraw.Draw(sheet)
    for i, (im, label) in enumerate(thumbs):
        r, c = divmod(i, 2)
        x, y = 8 + c * (cw + 8), 8 + r * (ch + 8)
        sheet.paste(im, (x + (cw - 24 - im.width) // 2, y))
        d.text((x + (cw - 24) // 2 - 90, y + im.height + 14), label, fill="#0E2841")
    sheet.save(ASSETS / "models_grid.png")
    return ASSETS / "s25_small.png", ASSETS / "models_grid.png"


# ---------------------------------------------------------------- замена картинки-контента
def swap_content_image(slide, new_img):
    # декор повторяется на многих слайдах — контентная картка уникальна
    seen = {}
    for sh in slide.shapes:
        if sh.shape_type == 13:
            h = hashlib.sha1(sh.image.blob).hexdigest()[:10]
            seen[h] = seen.get(h, 0) + 1
    # повторяющиеся блобы соберём позже по всем слайдам; здесь проще:
    # контент — самая большая картинка слайда
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    content = max(pics, key=lambda sh: sh.width * sh.height)
    left, top = content.left, content.top
    w, h = content.width, content.height
    content._element.getparent().remove(content._element)
    im = Image.open(new_img)
    aspect = im.width / im.height
    box_aspect = w / h
    if aspect > box_aspect:
        new_w, new_h = w, int(w / aspect)
    else:
        new_h, new_w = h, int(h * aspect)
    slide.shapes.add_picture(str(new_img), left + (w - new_w) // 2,
                             top + (h - new_h) // 2, width=new_w, height=new_h)


# ---------------------------------------------------------------- основная
def main(path: str):
    prs = Presentation(path)
    original_slides = list(prs.slides)   # ДО создания новых — только оригиналы
    pic_tpl = find_slide(prs, "Сквозная аналитика: слои")
    bul_tpl = find_slide(prs, "Как выбирать модель")

    s25_small, models_grid = prep_assets()

    # --- новые слайды вместо вставленных ---
    repl = {}
    repl[22] = make_pic_slide(prs, pic_tpl, "От каналов — к пути клиента",
        ASSETS / "s22_4.png",
        "Дальше — про то, что между кликом и деньгами: путь пользователя и атрибуция")
    repl[23] = make_pic_slide(prs, pic_tpl, "Пользовательский путь",
        ASSETS / "s23_4.png",
        "Все шаги от первого знакомства до покупки и лояльности — не один клик из одного канала")
    repl[24] = make_pic_slide(prs, pic_tpl, "Проблема поканальной аналитики",
        ASSETS / "s24_4.png",
        "Каждый канал рапортует о своих конверсиях: суммы не сходятся, бюджеты делятся вслепую")
    repl[25] = make_pic_slide(prs, pic_tpl, "Как это выглядит в кабинете",
        s25_small,
        "Цифры по каналам есть, ответа на вопрос «кто принёс продажу» — нет")
    repl[32] = make_pic_slide(prs, pic_tpl, "Каждый канал — сам себе молодец",
        ASSETS / "s32_4.png",
        "Расходы и конверсии по каналам раздельно: кому засчитать совместную конверсию?")
    repl[34] = make_pic_slide(prs, pic_tpl, "Какие пути бывают?",
        ASSETS / "s34_2052.png",
        "От одного касания до петель и повторов — модель атрибуции должна работать на любых")
    repl[35] = make_pic_slide(prs, pic_tpl, "Четыре правила на одном пути",
        models_grid,
        "Один и тот же путь — четыре способа поделить ценность: первому, последнему, поровну, с затуханием")
    repl[39] = make_bul_slide(prs, bul_tpl, "Атрибуция глазами диджитал-специалиста",
        ["Результат: продажи",
         "Усилия: онлайн- и офлайн-кампании, сайты",
         "Реклама в агрегаторах, CRM и т.д.",
         "Смотрит на весь путь, а не на канал"],
        left_head="Диджитал-специалист")
    repl[40] = make_bul_slide(prs, bul_tpl, "Атрибуция глазами e-commerce руководителя",
        ["Усилия: реклама, торговые предложения, витрина",
         "Результат: прибыль, клиенты, доля рынка",
         "«На какой канал добавить бюджет?»",
         "«Окупаются ли инвестиции в привлечение?»"],
        left_head="E-commerce руководитель")
    repl[41] = make_pic_slide(prs, pic_tpl, "Всех помирит — модель атрибуции",
        ASSETS / "s41_4.png",
        "Правило, какой переход считать источником визита: ручное или автоматическое (ML)")

    # --- «под капотом»: шаги 1–2 после пайплайна vc.ru, шаг 3 — замена картинки ---
    h1 = make_pic_slide(prs, pic_tpl, "Шаг 1 · цепочки → признаки",
        HOOD / "hood_01_features.png",
        "Счётчики касаний по каналам + длина пути; метка «купил» приходит из CRM (разметка из 3.2!)")
    h2 = make_pic_slide(prs, pic_tpl, "Шаг 2 · модель учится",
        HOOD / "hood_02_train.png",
        "Логрегрессия подбирает веса так, чтобы P росла у купивших: log-loss 0.69 → 0.37")
    swap_slide = find_slide(prs, "Вклад касания = разница прогнозов")
    swap_content_image(swap_slide, HOOD / "hood_03_example.png")

    # --- пересборка порядка ---
    drop_set = {22, 23, 24, 25, 32, 34, 36, 37, 38, 39, 40, 41}   # 35-й заменяется композитом
    replaced_by = {22: repl[22], 23: repl[23], 24: repl[24], 25: repl[25],
                   32: repl[32], 34: repl[34], 35: repl[35],
                   39: repl[39], 40: repl[40], 41: repl[41]}

    final = []
    for i, s in enumerate(original_slides, 1):
        if i in replaced_by:
            final.append(replaced_by[i])
            continue
        if i in drop_set:
            continue
        final.append(s)
        if slide_title(s).startswith("Пайплайн ML-атрибуции"):
            final.append(h1)
            final.append(h2)

    # sldId-элемент по части слайда (part стабилен между обходами, в отличие от прокси Slide)
    el_by_part = {}
    for el in list(prs.slides._sldIdLst):
        rId = el.get(qn("r:id"))
        part = prs.part.rels[rId].target_part
        el_by_part[id(part)] = el

    # удалить выброшенные оригиналы (rels + элементы)
    for i in sorted(drop_set):
        el = el_by_part[id(original_slides[i - 1].part)]
        prs.part.drop_rel(el.get(qn("r:id")))
        prs.slides._sldIdLst.remove(el)

    # выстроить финальный порядок
    sldIdLst = prs.slides._sldIdLst
    for el in list(sldIdLst):
        sldIdLst.remove(el)
    for s in final:
        sldIdLst.append(el_by_part[id(s.part)])

    renumber(prs)
    prs.save(path)
    print("готово:", path, "| слайдов:", len(list(prs.slides)))


if __name__ == "__main__":
    main(sys.argv[1])
