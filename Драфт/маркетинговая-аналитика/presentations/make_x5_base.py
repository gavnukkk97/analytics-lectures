# -*- coding: utf-8 -*-
"""Сборка компактной базы носителей из корпоративного шаблона X5Digital.

Готовые X5-презентации содержат корпоративную графику и в открытый репозиторий
не выкладываются. Чтобы пересобрать их на своей машине:

1. Положите файл шаблона (X5Digital_Шаблон_презентации_*.pptx) рядом со скриптом
   или укажите путь первым аргументом.
2. python3 presentations/make_x5_base.py [путь/к/шаблону.pptx]
   → рядом появится x5_base.pptx (4 слайда-носителя, ~5 МБ).
3. python3 presentations/make_decks_x5.py — генерирует 21 презентацию в папку
   presentations_x5/ (рядом с базой).
"""
from __future__ import annotations
import io
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from PIL import Image

HERE = Path(__file__).resolve().parent
KEEP = [1, 3, 10, 11]  # титул; заголовок+пояснение; текст и список; важное число


def build(template_path: Path, out_path: Path) -> None:
    prs = Presentation(str(template_path))
    sldIds = list(prs.slides._sldIdLst)
    for idx in range(len(sldIds) - 1, -1, -1):
        if (idx + 1) not in KEEP:
            prs.part.drop_rel(sldIds[idx].get(qn("r:id")))
            prs.slides._sldIdLst.remove(sldIds[idx])
    raw = out_path.with_suffix(".raw.pptx")
    prs.save(str(raw))

    # пережимаем крупные PNG (>400 КБ) до 1600px
    zin = zipfile.ZipFile(str(raw))
    with zipfile.ZipFile(str(out_path), "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if (item.filename.startswith("ppt/media/") and item.file_size > 400_000
                    and item.filename.endswith(".png")):
                img = Image.open(io.BytesIO(data))
                if img.width > 1600:
                    img = img.resize((1600, int(img.height * 1600 / img.width)), Image.LANCZOS)
                buf = io.BytesIO()
                img.save(buf, "PNG", optimize=True)
                if buf.tell() < len(data):
                    data = buf.getvalue()
            zout.writestr(item, data)
    raw.unlink()
    print("база:", out_path, f"({out_path.stat().st_size / 1e6:.1f} МБ)")


if __name__ == "__main__":
    tpl = Path(sys.argv[1]) if len(sys.argv) > 1 else next(
        (p for p in HERE.glob("X5Digital*.pptx")), None)
    if tpl is None or not tpl.exists():
        sys.exit("Шаблон X5Digital не найден: положите файл рядом или передайте путь аргументом.")
    build(tpl, HERE / "x5_base.pptx")
